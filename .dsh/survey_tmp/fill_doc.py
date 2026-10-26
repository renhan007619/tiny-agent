# -*- coding: utf-8 -*-
"""将验证过的数据填入 report.docx -> filled_report.docx（按模板真实合并结构）"""
import gen_data as G
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

CEN = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT
FONT = "宋体"

def style_run(run, size=9):
    run.font.name = FONT
    run.font.size = Pt(size)
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:eastAsia"), FONT)

def set_cell(cell, text, align=CEN, size=9):
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = align
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        if i > 0:
            br = p.add_run()
            br.add_break()
            style_run(br, size)
        run = p.add_run(ln)
        style_run(run, size)

def shrink_margins(cell, twips=20):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("w:left", "w:right"):
        el = OxmlElement(side)
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)

def sd(v, nd=1):
    if v == 0:
        return "0"
    return f"{v:+.{nd}f}" if nd else f"{v:+d}"

doc = docx.Document("report.docx")

# ============ 实验一 普通水准 T0：逐点连续记录（转点单元格骑跨测站行） ============
T0 = doc.tables[0]
# 测站号：{r2,r3}=1 {r4,r5}=2 {r6,r8}=3 {r9,r10}=4 {r11,r12}=5（写在合并起始行）
for r, v in [(2, 1), (4, 2), (6, 3), (9, 4), (11, 5)]:
    set_cell(T0.rows[r].cells[0], str(v))
# 点号/高程：起点行r2，各转点 {r3,r4}{r5,r6}{r7,r9}{r10,r11}，终点行{r12,r13}
points = [(2, "BM.A", 156.328), (3, "P1", 156.646), (5, "P2", 156.501),
          (7, "P3", 156.099), (10, "P4", 156.386), (12, "BM.A", 156.336)]
for r, name, H in points:
    set_cell(T0.rows[r].cells[2], name)
    set_cell(T0.rows[r].cells[7], f"{H:.3f}")
# 后视a：各后视点行；前视b：各前视点行（与点号同格位）
backs = [(2, 1.647), (3, 1.283), (5, 0.984), (7, 1.516), (10, 1.371)]
for r, a in backs:
    set_cell(T0.rows[r].cells[3], f"{a:.3f}")
fronts = [(3, 1.329), (5, 1.428), (7, 1.386), (10, 1.229), (12, 1.421)]
for r, b in fronts:
    set_cell(T0.rows[r].cells[4], f"{b:.3f}")
# 高差：正{r2,r3} 负{r4,r5} 负{r6,r7} 正{r8,r10} 负{r11,r12}
for r, col, h in [(2, 5, 0.318), (4, 6, 0.145), (6, 6, 0.402), (8, 5, 0.287), (11, 6, 0.050)]:
    set_cell(T0.rows[r].cells[col], f"{h:.3f}")
# 备注（整个数据区一个大单元格）
set_cell(T0.rows[2].cells[8], "H_BM.A=156.328m（已知）；fh=Σh=+8mm≤±12√5≈±27mm")
# Σ行 / 检核行
set_cell(T0.rows[33].cells[3], f"{G.sum_a:.3f}")
set_cell(T0.rows[33].cells[4], f"{G.sum_b:.3f}")
set_cell(T0.rows[33].cells[5], f"{G.sum_pos:.3f}")
set_cell(T0.rows[33].cells[6], f"{G.sum_neg:.3f}")
set_cell(T0.rows[34].cells[3],
         f"Σa={G.sum_a:.3f}m  Σb={G.sum_b:.3f}m  Σh={G.sum_h:+.3f}m\n"
         f"Σa-Σb={G.sum_a - G.sum_b:+.3f}m，计算检核相符")

# ============ 实验二 全站仪 T1：每2行一个观测单元（测点/水平角/备注跨2行） ============
T1 = doc.tables[1]
for k, (beta, rows, note) in enumerate(G.exp2_calc):
    rA, rB = 2 * k + 1, 2 * k + 2
    set_cell(T1.rows[rA].cells[0], "O")
    set_cell(T1.rows[rA].cells[3], G.fmt_sec(beta), size=8)
    a, b = rows
    set_cell(T1.rows[rA].cells[1], "A")
    set_cell(T1.rows[rA].cells[2], G.dms(*a["hr"]), size=8)
    set_cell(T1.rows[rA].cells[4], G.vdms(*a["va"]), size=8)
    set_cell(T1.rows[rA].cells[5], f"{a['sd']:.3f}")
    set_cell(T1.rows[rA].cells[6], f"{a['hd']:.3f}")
    set_cell(T1.rows[rA].cells[7], f"{a['vd']:+.3f}")
    set_cell(T1.rows[rB].cells[1], "B")
    set_cell(T1.rows[rB].cells[2], G.dms(*b["hr"]), size=8)
    set_cell(T1.rows[rB].cells[4], G.vdms(*b["va"]), size=8)
    set_cell(T1.rows[rB].cells[5], f"{b['sd']:.3f}")
    set_cell(T1.rows[rB].cells[6], f"{b['hd']:.3f}")
    set_cell(T1.rows[rB].cells[7], f"{b['vd']:+.3f}")
    if note:
        set_cell(T1.rows[rA].cells[8], note)

# ============ 实验三 T2 水平角 / T3 竖直角（原逻辑已核对无误） ============
T2 = doc.tables[2]
hb = G.half_beta
rows2 = [
    (1, {"c0": "O", "c1": "1", "c3": "A", "c4": G.dms(*G.exp3h[0][3])}),
    (2, {"c3": "B", "c4": G.dms(*G.exp3h[1][3]), "c5": G.fmt_sec(hb[(1, "左")])}),
    (3, {"c3": "B", "c4": G.dms(*G.exp3h[2][3])}),
    (4, {"c3": "A", "c4": G.dms(*G.exp3h[3][3]), "c5": G.fmt_sec(hb[(1, "右")]), "c6": G.fmt_sec(G.beta1)}),
    (5, {"c0": "O", "c1": "2", "c3": "A", "c4": G.dms(*G.exp3h[4][3])}),
    (6, {"c3": "B", "c4": G.dms(*G.exp3h[5][3]), "c5": G.fmt_sec(hb[(2, "左")])}),
    (7, {"c3": "B", "c4": G.dms(*G.exp3h[6][3])}),
    (8, {"c3": "A", "c4": G.dms(*G.exp3h[7][3]), "c5": G.fmt_sec(hb[(2, "右")]), "c6": G.fmt_sec(G.beta2), "c7": G.fmt_sec(G.beta_avg)}),
]
for r, vals in rows2:
    for k, v in vals.items():
        set_cell(T2.rows[r].cells[int(k[1:])], v, size=8 if int(k[1:]) >= 4 else 9)

T3 = doc.tables[3]
seq = [("A", 1), ("A", 2), ("B", 1), ("B", 2)]
for k, (tgt, setn) in enumerate(seq):
    rL, rR = 1 + 2 * k, 2 + 2 * k
    al, ar, x, am = G.v_res[(tgt, setn)]
    Lread = next(r for t, s, p, r in G.exp3v if (t, s, p) == (tgt, setn, "左"))
    Rread = next(r for t, s, p, r in G.exp3v if (t, s, p) == (tgt, setn, "右"))
    cL = T3.rows[rL].cells
    if k % 2 == 0:
        set_cell(cL[0], "O")
    set_cell(cL[1], tgt)
    set_cell(cL[3], G.dms(*Lread), size=8)
    shrink_margins(cL[3])
    set_cell(cL[4], G.fmt_sec(al), size=8)
    cR = T3.rows[rR].cells
    set_cell(cR[3], G.dms(*Rread), size=8)
    shrink_margins(cR[3])
    set_cell(cR[4], G.fmt_sec(ar), size=8)
    set_cell(cR[5], f"{x:+.0f}″")
    set_cell(cR[6], G.fmt_sec(am), size=8)
    if k % 2 == 1:
        set_cell(cR[7], G.fmt_sec(G.v_avg[tgt]), size=8)

# ============ 实验四 T4 四等水准：黑面/红面为后+前合并格（写两行） ============
T4 = doc.tables[4]
for i, s in enumerate(G.st4):
    A, B, C, D = (T4.rows[4 + 4 * i + j].cells for j in range(4))
    set_cell(A[0], str(i + 1))                       # 测站编号（跨4行）
    set_cell(A[1], f"{s['bk']}→{s['fr']}")           # 点号（跨4行）
    set_cell(A[2], f"{s['bl']:.3f}")                 # 后尺下丝
    set_cell(A[4], f"{s['fl']:.3f}")                 # 前尺下丝
    set_cell(A[6], f"后{s['no_b']}")
    set_cell(A[7], f"{s['a']:.3f}\n{s['b']:.3f}")    # 黑面：后黑/前黑
    set_cell(A[8], f"{s['a_red']:.3f}\n{s['b_red']:.3f}")
    set_cell(A[9], sd(s["ea"], 0))
    set_cell(B[2], f"{s['bu']:.3f}")                 # 后尺上丝
    set_cell(B[4], f"{s['fu']:.3f}")                 # 前尺上丝
    set_cell(B[6], f"前{s['no_f']}")
    set_cell(B[9], sd(s["eb"], 0))
    set_cell(C[2], f"{s['Db']:.1f}")                 # 后视距
    set_cell(C[4], f"{s['Df']:.1f}")                 # 前视距
    set_cell(C[7], sd(s["h_b"], 3))
    set_cell(C[8], sd(s["h_r"], 3))
    set_cell(C[9], sd(s["ediff"], 0))
    set_cell(C[10], sd(s["h_mid"], 4))               # 高差中数（跨4行）
    set_cell(D[2], sd(s["d"]))
    set_cell(D[4], sd(s["cum"]))
set_cell(T4.rows[4].cells[11], "K1=4.687m\nK2=4.787m")   # 备注（跨站1-2）
sum_Db = round(sum(s["Db"] for s in G.st4), 1)
sum_Df = round(sum(s["Df"] for s in G.st4), 1)
set_cell(T4.rows[28].cells[1],
         f"Σ后视距={sum_Db:.1f}m  Σ前视距={sum_Df:.1f}m  总视距={G.L4:.1f}m\n"
         f"Σh黑={G.sum_hb:+.3f}m  Σh红={G.sum_hr:+.3f}m  Σh中={G.sum_hm:+.4f}m，检核相符")

# ============ 实验四 T5 成果表：阶梯式（段值{奇,偶}行 / 点号{偶,奇}行） ============
T5 = doc.tables[5]
set_cell(T5.rows[1].cells[0], "BM.A")
set_cell(T5.rows[1].cells[6], f"{G.H_A4:.3f}")
for j, (s, v) in enumerate(zip(G.st4, G.v4)):
    r_seg = 2 * j + 1          # 段值单元格 {r_seg, r_seg+1}
    r_pt = 2 * j + 2           # 到达点单元格 {r_pt, r_pt+1}
    set_cell(T5.rows[r_seg].cells[1], f"{s['Db'] + s['Df']:.1f}")
    set_cell(T5.rows[r_seg].cells[2], "1")
    set_cell(T5.rows[r_seg].cells[3], sd(s["h_mid"], 4))
    set_cell(T5.rows[r_seg].cells[4], sd(v, 0))
    set_cell(T5.rows[r_seg].cells[5], sd(round(s["h_mid"] + v / 1000, 4), 4))
    set_cell(T5.rows[r_pt].cells[0], s["fr"])
    set_cell(T5.rows[r_pt].cells[6], f"{G.H4[s['fr']]:.3f}")
c22 = T5.rows[22].cells
set_cell(c22[1], f"{G.L4:.1f}")
set_cell(c22[2], "6")
set_cell(c22[3], sd(G.sum_hm, 4))
set_cell(c22[4], sd(-G.fh4_mm, 0))
set_cell(c22[5], "0.000")
set_cell(T5.rows[1].cells[7],
         f"fh=Σh测={G.sum_hm:+.4f}m=+{G.fh4_mm}mm\n"
         f"fh容=±20√L=±20×√{G.L4 / 1000:.3f}=±{G.tol4_mm:.1f}mm\n"
         f"|fh|≤|fh容|，成果合格", align=LEFT)

# ============ 页眉信息 ============
dates = [("2026年10月13日", "晴"), ("2026年10月16日", "多云"),
         ("2026年10月20日", "阴"), ("2026年10月23日", "晴")]
di = 0
for p in doc.paragraphs:
    t = p.text
    if t.startswith("实验日期") and di < 4:
        p.runs[1].text = dates[di][0]
        p.runs[4].text = dates[di][1]
        di += 1
    elif t.startswith("学年学期"):
        p.runs[1].text = " 2026—2027学年第一学期"
    elif t.startswith("仪器") and "班组" in t:
        p.runs[1].text = " DS3水准仪"
        p.runs[5].text = " 晴"
print("dates filled:", di)

# ============ 路线草图 ============
import sketch
sketch.make("sketch1.png", nodes=[("BM.A", 0.5, 1.9, "H=156.328m"), ("P1", 2.3, 1.5, ""),
                                  ("P2", 2.7, 0.3, ""), ("P3", 1.3, -0.5, ""), ("P4", -0.3, 0.2, "")])
sketch.make("sketch4.png", nodes=[("BM.A", -0.3, 1.7, "H=162.474m"), ("P1", 1.2, 2.0, ""),
                                  ("P2", 2.6, 1.0, ""), ("P3", 2.2, -0.5, ""),
                                  ("P4", 0.5, -1.0, ""), ("P5", -1.0, -0.3, "")])

def new_para_after(anchor):
    el = OxmlElement("w:p")
    anchor._p.addnext(el)
    return Paragraph(el, anchor._parent)

anchors = [p for p in doc.paragraphs if "水准测量路线草图" in p.text]
caps = ["（闭合水准路线，全长约320 m，共5站）", "（闭合水准路线，全长约608 m，共6站）"]
pngs = ["sketch1.png", "sketch4.png"]
for anc, png, cap in zip(anchors, pngs, caps):
    p1 = new_para_after(anc)
    p1.alignment = CEN
    p1.add_run().add_picture(png, width=Cm(9.5))
    p2 = new_para_after(p1)
    p2.alignment = CEN
    r = p2.add_run(cap)
    style_run(r, 10.5)

doc.save("filled_report.docx")
print("saved filled_report.docx")
