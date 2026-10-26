# -*- coding: utf-8 -*-
"""从 filled_report.docx 反向读取数据，独立验证自洽性（按模板真实合并结构）"""
import re
import math
import docx

d = docx.Document("filled_report.docx")
ok = [0]

def chk(cond, msg):
    if not cond:
        raise AssertionError(msg)
    ok[0] += 1

def f(x):
    return float(x.replace("+", ""))

def psec(s):
    m = re.match(r"(-?)(\d+)°(\d+)′(\d+)″", s)
    sg = -1 if m.group(1) else 1
    return sg * (int(m.group(2)) * 3600 + int(m.group(3)) * 60 + int(m.group(4)))

# ---------- T0 普通水准（逐点连续记录） ----------
t = d.tables[0]
# 点行: r2=BM.A(起点) r3=P1 r5=P2 r7=P3 r10=P4 r12=BM.A(终点)
prow = {"BM.A_start": 2, "P1": 3, "P2": 5, "P3": 7, "P4": 10, "BM.A_end": 12}
H = 156.328
sum_a = sum_b = sum_h = 0.0
seq = ["P1", "P2", "P3", "P4", "BM.A_end"]
for i, pt in enumerate(seq):
    prev = "BM.A_start" if i == 0 else seq[i - 1]
    rp, rc = prow[prev], prow[pt]
    a = f(t.rows[rp].cells[3].text)
    b = f(t.rows[rc].cells[4].text)
    h = a - b
    # 高差写在 {r_start..} 单元格：h1=(2,5) h2=(4,6) h3=(6,6) h4=(8,5) h5=(11,6)，col5=正 col6=负(绝对值)
    hcell = [(2, 5), (4, 6), (6, 6), (8, 5), (11, 6)][i]
    hv = f(t.rows[hcell[0]].cells[hcell[1]].text)
    hv_signed = hv if hcell[1] == 5 else -hv
    chk(abs(h - hv_signed) < 1e-9, f"T0 站{i+1} 高差列与a-b不符")
    Hp = f(t.rows[rp].cells[7].text)
    Hc = f(t.rows[rc].cells[7].text)
    chk(abs(Hp - H) < 1e-9, f"T0 站{i+1} 后点高程链断")
    H = round(H + h, 3)
    chk(abs(Hc - H) < 1e-9, f"T0 站{i+1} 前点高程不符")
    sum_a += a; sum_b += b; sum_h += h
chk(abs(sum_h - 0.008) < 1e-9, "T0 Σh")
chk(abs(sum_h) <= 12 * math.sqrt(5) / 1000, "T0 闭合差超限")
chk(f(t.rows[33].cells[3].text) == round(sum_a, 3), "T0 Σa")
chk(f(t.rows[33].cells[4].text) == round(sum_b, 3), "T0 Σb")
print(f"T0 ok: Σa={sum_a:.3f} Σb={sum_b:.3f} fh={sum_h*1000:+.0f}mm 终点高程={H:.3f}")

# ---------- T1 全站仪（3单元×2目标） ----------
t = d.tables[1]
betas = []
for k in range(3):
    rA, rB = 2 * k + 1, 2 * k + 2
    cA, cB = t.rows[rA].cells, t.rows[rB].cells
    chk(cA[0].text.strip() == "O" and cB[0].text.strip() == "O", "T1 测点")
    chk(cA[1].text.strip() == "A" and cB[1].text.strip() == "B", "T1 目标")
    beta = psec(cB[2].text) - psec(cA[2].text)
    chk(abs(beta - psec(cA[3].text)) < 1e-6, f"T1 单元{k+1} 水平角≠读数差")
    betas.append(beta)
    for c in (cA, cB):
        sd_ = f(c[5].text); hd = f(c[6].text); vd = f(c[7].text)
        va = psec(c[4].text)
        chk(abs(math.hypot(hd, vd) - sd_) < 0.003, "T1 SD/HD/VD")
        chk(abs(hd - sd_ * math.cos(math.radians(va / 3600))) < 0.0015, "T1 HD-竖直角")
chk(max(betas) - min(betas) <= 40, "T1 水平角互差")
print(f"T1 ok: 水平角 {[f'{b/3600:.6f}°' for b in betas]} 互差{(max(betas)-min(betas)):.0f}″")

# ---------- T2 水平角 ----------
t = d.tables[2]
rd = [psec(t.rows[r].cells[4].text) for r in range(1, 9)]
for s_, (i1, i2, i3, i4) in [(1, (0, 1, 2, 3)), (2, (4, 5, 6, 7))]:
    bl = rd[i2] - rd[i1]; br = rd[i3] - rd[i4]
    chk(abs(bl - br) <= 40, f"T2 测回{s_} 互差>40″")
    beta = (bl + br) / 2
    chk(abs(beta - psec(t.rows[4 if s_ == 1 else 8].cells[6].text)) < 1e-6, f"T2 测回{s_} 一测回")
    chk(abs(bl - psec(t.rows[2 if s_ == 1 else 6].cells[5].text)) < 1e-6, f"T2 测回{s_} 半测回左")
    chk(abs(br - psec(t.rows[4 if s_ == 1 else 8].cells[5].text)) < 1e-6, f"T2 测回{s_} 半测回右")
b1 = psec(t.rows[4].cells[6].text); b2 = psec(t.rows[8].cells[6].text)
chk(abs((b1 + b2) / 2 - psec(t.rows[8].cells[7].text)) <= 1, "T2 平均角值")
chk(abs(b1 - b2) <= 40, "T2 测回间互差")
print("T2 ok: 水平角全部自洽")

# ---------- T3 竖直角 ----------
t = d.tables[3]
for k, tgt in enumerate(["A", "A", "B", "B"]):
    rL, rR = 1 + 2 * k, 2 + 2 * k
    L = psec(t.rows[rL].cells[3].text); R = psec(t.rows[rR].cells[3].text)
    al = 90 * 3600 - L; ar = R - 270 * 3600
    x = (L + R - 360 * 3600) / 2
    chk(abs(al - psec(t.rows[rL].cells[4].text)) < 1e-6, f"T3 {tgt} α左")
    chk(abs(ar - psec(t.rows[rR].cells[4].text)) < 1e-6, f"T3 {tgt} α右")
    chk(abs(x - f(t.rows[rR].cells[5].text.replace("″", ""))) < 1e-6, f"T3 {tgt} 指标差")
    chk(abs(x) <= 25, f"T3 {tgt} 指标差超限")
    am = (al + ar) / 2
    chk(abs(am - psec(t.rows[rR].cells[6].text)) < 1e-6, f"T3 {tgt} 一测回")
    if k % 2 == 1:
        other = psec(t.rows[rR - 2].cells[6].text)
        chk(abs((am + other) / 2 - psec(t.rows[rR].cells[7].text)) <= 1, f"T3 {tgt} 平均")
print("T3 ok: 竖直角/指标差全部自洽")

# ---------- T4 四等水准 ----------
t = d.tables[4]
K = {1: 4.687, 2: 4.787}
cum = 0.0
tot_b = tot_f = tot_hb = tot_hr = tot_hm = 0.0
for i in range(6):
    A, B, C, D = (t.rows[4 + 4 * i + j].cells for j in range(4))
    no_b = int(A[6].text.strip()[1]); no_f = int(B[6].text.strip()[1])
    bl, bu = f(A[2].text), f(B[2].text)
    fl, fu = f(A[4].text), f(B[4].text)
    abl, bbl = A[7].text.split("\n")          # 黑面：后黑/前黑
    ar_, br_ = A[8].text.split("\n")          # 红面：后红/前红
    a, b = f(abl), f(bbl)
    a_r, b_r = f(ar_), f(br_)
    ea = f(A[9].text) if A[9].text.strip() != "0" else 0.0
    eb = f(B[9].text) if B[9].text.strip() != "0" else 0.0
    Db, Df = f(C[2].text), f(C[4].text)
    h_b, h_r = f(C[7].text), f(C[8].text)
    ed = f(C[9].text) if C[9].text.strip() != "0" else 0.0
    h_m = f(C[10].text)
    chk(abs((bl - bu) * 100 - Db) < 1e-9, f"T4 站{i+1} 后视距")
    chk(abs((fl - fu) * 100 - Df) < 1e-9, f"T4 站{i+1} 前视距")
    chk(abs((a + K[no_b] - a_r) * 1000 - ea) < 1e-6, f"T4 站{i+1} 后K检核")
    chk(abs((b + K[no_f] - b_r) * 1000 - eb) < 1e-6, f"T4 站{i+1} 前K检核")
    chk(abs(ea) <= 3 and abs(eb) <= 3, f"T4 站{i+1} 黑红读数差>3mm")
    chk(abs((a - b) - h_b) < 1e-9, f"T4 站{i+1} h黑")
    chk(abs((a_r - b_r) - h_r) < 1e-9, f"T4 站{i+1} h红")
    chk(abs(h_b - (h_r - (K[no_b] - K[no_f]) + ed / 1000)) < 2e-9, f"T4 站{i+1} 黑红高差关系")
    chk(abs(ed) <= 5, f"T4 站{i+1} 黑红高差之差>5mm")
    chk(abs((h_b - ed / 2000) - h_m) < 1e-9, f"T4 站{i+1} 高差中数")
    dcur = f(D[2].text) if D[2].text.strip() != "0" else 0.0
    chk(abs(Db - Df - dcur) < 1e-9, f"T4 站{i+1} 视距差")
    chk(abs(dcur) <= 3, f"T4 站{i+1} 视距差>3m")
    cum = round(cum + dcur, 1)
    chk(abs(cum - (f(D[4].text) if D[4].text.strip() != "0" else 0.0)) < 1e-9, f"T4 站{i+1} 累积视距差")
    chk(abs(cum) <= 10, "T4 累积视距差>10m")
    tot_b += Db; tot_f += Df; tot_hb += h_b; tot_hr += h_r; tot_hm += h_m
fh = round(tot_hm, 4)
chk(abs((tot_hb + tot_hr) / 2 - fh) < 1e-9, "T4 Σh中检核")
chk(abs(fh) <= 0.020 * math.sqrt((tot_b + tot_f) / 1000), "T4 闭合差超限")
print(f"T4 ok: Σ后距={tot_b:.1f} Σ前距={tot_f:.1f} fh={fh*1000:+.0f}mm")

# ---------- T5 成果表（阶梯式） ----------
t = d.tables[5]
chk(t.rows[1].cells[0].text.strip() == "BM.A", "T5 起点")
chk(f(t.rows[1].cells[6].text) == 162.474, "T5 起点高程")
H = 162.474
sv = 0
for j in range(6):
    r_seg, r_pt = 2 * j + 1, 2 * j + 2
    c = t.rows[r_seg].cells
    L = f(c[1].text); hm = f(c[3].text); v = f(c[4].text) if c[4].text.strip() != "0" else 0.0
    hcorr = f(c[5].text)
    chk(abs(hm + v / 1000 - hcorr) < 1e-9, f"T5 段{j+1} 改正后高差")
    chk(90 <= L <= 115, f"T5 段{j+1} 距离异常")
    H = round(H + hcorr, 4)
    pc = t.rows[r_pt].cells
    chk(abs(H - f(pc[6].text)) < 1e-9, f"T5 段{j+1} 高程链")
    sv += v
chk(round(H, 3) == 162.474, "T5 不闭合")
chk(sv == -9, "T5 改正数总和")
c22 = t.rows[22].cells
chk(f(c22[1].text) == round(tot_b + tot_f, 1), "T5 Σ距离")
chk(f(c22[3].text) == round(tot_hm, 4), "T5 Σ实测高差")
print(f"T5 ok: 平差后闭合 {H:.4f}, Σv={sv:+.0f}mm")
print(f"ALL {ok[0]} ROUND-TRIP CHECKS PASSED")
